"""Build the Game Night Google Sheet as an .xlsx (uploaded to Drive, which converts it).
Usage: python build_sheet.py out.xlsx [--sample]   (--sample fills test data for a local check)"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment

out = sys.argv[1]
SAMPLE = "--sample" in sys.argv
NP, NB, NM = 25, 200, 80          # player rows, buy-in rows, match rows
P_END, B_END, M_END = NP + 1, NB + 1, NM + 1

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F4E3D")
SEC_FILL = PatternFill("solid", fgColor="DCE8E2")
IN_FILL = PatternFill("solid", fgColor="FFF7CC")
LIVE_FILL = PatternFill("solid", fgColor="FBD9D5")
GOLD_FILL = PatternFill("solid", fgColor="FBEFC9")
thin = Side(style="thin", color="D0D7D3")
BORDER = Border(bottom=thin)
MONEY = '$#,##0.00;-$#,##0.00;"-"'
MONEY0 = '$#,##0.00'

wb = Workbook()

def style_header(ws, row, ncols, start=1):
    for c in range(start, start + ncols):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def base_font(ws, max_row, max_col):
    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for cell in row:
            if cell.font is None or cell.font.name != FONT:
                cell.font = Font(name=FONT, bold=cell.font.bold if cell.font else False,
                                 color=cell.font.color if cell.font else None,
                                 size=cell.font.size if cell.font else 11)

def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w

def fill_range(ws, rng, fill):
    ws.conditional_formatting.add(rng, FormulaRule(formula=['TRUE'], fill=fill))

# ---------------- Start here ----------------
st = wb.active
st.title = "Start here"
lines = [
    ("Game Night", "title"),
    ("Poker buy-ins and the FIFA tournament, shared by everyone with the link.", None),
    ("", None),
    ("1. Players: add your name, pick your group, usual team, and whether you play poker and FIFA.", None),
    ("2. Buy-ins: add one row each time someone buys in. Leave Amount blank for a standard $20 buy-in.", None),
    ("3. Matches: when a FIFA game starts, add a row with both players and teams and set Status to Live.", None),
    ("    Update the goals as you play. Set Status to Final when it ends. For a knockout draw, pick who won on penalties.", None),
    ("4. Cash-out: at the end of the night, type how many chips of each colour each player has.", None),
    ("5. Leaderboard: live matches, group tables, the overall FIFA table and poker results update by themselves.", None),
    ("", None),
    ("Yellow cells are the ones to fill in. Everything else is calculated.", None),
    ("Group names, the buy-in amount and chip values are on the Settings tab.", None),
    ("Only Final matches count. Friendlies never count. 3 points for a win, 1 for a draw; ties split by goal difference, then goals scored.", None),
]
for i, (text, kind) in enumerate(lines, start=1):
    c = st.cell(row=i, column=1, value=text)
    c.font = Font(name=FONT, size=20 if kind == "title" else 11, bold=kind == "title", color="1F4E3D" if kind == "title" else None)
st.column_dimensions["A"].width = 120

# ---------------- Settings ----------------
se = wb.create_sheet("Settings")
se["A1"] = "Settings"; se["A1"].font = Font(name=FONT, size=16, bold=True)
se["A3"] = "Standard buy-in"; se["B3"] = 20
se["B3"].number_format = MONEY0; se["B3"].fill = IN_FILL; se["B3"].font = Font(name=FONT, color="0000FF")
se["A5"], se["B5"], se["C5"], se["D5"] = "Chip colour", "Value", "Chips per buy-in", "Value per buy-in"
style_header(se, 5, 4)
chips = [("White", 0.10, 10), ("Red", 0.20, 10), ("Green", 0.50, 4), ("Purple", 1, 10), ("Blue", 5, 1)]
for i, (n, v, k) in enumerate(chips):
    r = 6 + i
    se.cell(row=r, column=1, value=n)
    se.cell(row=r, column=2, value=v).number_format = MONEY0
    se.cell(row=r, column=3, value=k)
    se.cell(row=r, column=4, value=f"=B{r}*C{r}").number_format = MONEY0
    for col in (2, 3):
        se.cell(row=r, column=col).fill = IN_FILL
        se.cell(row=r, column=col).font = Font(name=FONT, color="0000FF")
se["A11"] = "Starting stack"; se["A11"].font = Font(name=FONT, bold=True)
se["C11"] = "=SUM(C6:C10)"; se["D11"] = "=SUM(D6:D10)"; se["D11"].number_format = MONEY0
se["C11"].font = se["D11"].font = Font(name=FONT, bold=True)
se["A12"] = '=IF(ABS(D11-B3)<0.005,"The stack matches the buy-in.","The stack is worth "&TEXT(D11,"$0.00")&" but the buy-in is "&TEXT(B3,"$0.00")&".")'
se["A14"] = "Groups (type the names)"; style_header(se, 14, 1)
fill_range(se, "A15:A22", IN_FILL)
se["A13"] = "Source: chip values and stack from the host, 17 Oct 2026."
se["A13"].font = Font(name=FONT, size=9, italic=True, color="5B6662")
widths(se, {"A": 28, "B": 12, "C": 18, "D": 18})

GROUPS = "Settings!$A$15:$A$22"

# ---------------- Players ----------------
pl = wb.create_sheet("Players")
for i, h in enumerate(["Name", "Group", "Usual team", "Poker", "FIFA"], start=1):
    pl.cell(row=1, column=i, value=h)
style_header(pl, 1, 5)
fill_range(pl, f"A2:E{P_END}", IN_FILL)
pl["A1"].comment = Comment("Type each name once, the same way every time.", "Game Night")
dv_group = DataValidation(type="list", formula1=f"={GROUPS}", allow_blank=True)
dv_yes = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
pl.add_data_validation(dv_group); pl.add_data_validation(dv_yes)
dv_group.add(f"B2:B{P_END}"); dv_yes.add(f"D2:E{P_END}")
widths(pl, {"A": 20, "B": 14, "C": 20, "D": 9, "E": 9})
pl.freeze_panes = "A2"

PNAMES = f"Players!$A$2:$A${P_END}"

# ---------------- Buy-ins ----------------
bi = wb.create_sheet("Buy-ins")
for i, h in enumerate(["Time", "Player", "Amount", "Recorded by"], start=1):
    bi.cell(row=1, column=i, value=h)
style_header(bi, 1, 4)
fill_range(bi, f"A2:D{B_END}", IN_FILL)
bi["C1"].comment = Comment("Leave blank for a standard buy-in (Settings). Type an amount only for a custom buy-in.", "Game Night")
bi["A1"].comment = Comment("Optional, e.g. 9:15 PM", "Game Night")
dv_pb = DataValidation(type="list", formula1=f"={PNAMES}", allow_blank=True)
bi.add_data_validation(dv_pb); dv_pb.add(f"B2:B{B_END}"); dv_pb.add(f"D2:D{B_END}")
for r in range(2, B_END + 1):
    bi.cell(row=r, column=3).number_format = MONEY0
widths(bi, {"A": 11, "B": 20, "C": 11, "D": 18})
bi.freeze_panes = "A2"

# ---------------- Matches ----------------
ma = wb.create_sheet("Matches")
mh = ["Stage", "Round", "Player 1", "Team 1", "Goals 1", "Goals 2", "Team 2", "Player 2", "Status", "Won on penalties", "Winner", "Group", "Recorded by", "Live order"]
for i, h in enumerate(mh, start=1):
    ma.cell(row=1, column=i, value=h)
style_header(ma, 1, len(mh))
fill_range(ma, f"A2:J{M_END}", IN_FILL); fill_range(ma, f"M2:M{M_END}", IN_FILL)
ma["A1"].comment = Comment("Group counts in the group table (both players must be in the same group). Knockout counts overall. Friendly never counts.", "Game Night")
ma["I1"].comment = Comment("Live while playing, Final when it ends. Only Final matches count.", "Game Night")
ma["J1"].comment = Comment("Only for a knockout that ends level.", "Game Night")
PG = f"Players!$A$2:$B${P_END}"
for r in range(2, M_END + 1):
    ma.cell(row=r, column=11, value=f'=IF(OR(C{r}="",H{r}="",I{r}<>"Final"),"",IF(N(E{r})>N(F{r}),C{r},IF(N(F{r})>N(E{r}),H{r},IF(J{r}<>"",J{r},"Draw"))))')
    ma.cell(row=r, column=12, value=f'=IFERROR(IF(VLOOKUP(C{r},{PG},2,0)=VLOOKUP(H{r},{PG},2,0),VLOOKUP(C{r},{PG},2,0)&"",""),"")')
    ma.cell(row=r, column=14, value=f'=IF(I{r}="Live",ROW(),"")')
dv_stage = DataValidation(type="list", formula1='"Group,Knockout,Friendly"', allow_blank=True)
dv_round = DataValidation(type="list", formula1='"Quarter-final,Semi-final,3rd place,Final"', allow_blank=True)
dv_pm = DataValidation(type="list", formula1=f"={PNAMES}", allow_blank=True)
dv_status = DataValidation(type="list", formula1='"Live,Final"', allow_blank=True)
dv_goals = DataValidation(type="whole", operator="between", formula1="0", formula2="50", allow_blank=True)
for dv in (dv_stage, dv_round, dv_pm, dv_status, dv_goals):
    ma.add_data_validation(dv)
dv_stage.add(f"A2:A{M_END}"); dv_round.add(f"B2:B{M_END}")
dv_pm.add(f"C2:C{M_END}"); dv_pm.add(f"H2:H{M_END}"); dv_pm.add(f"J2:J{M_END}"); dv_pm.add(f"M2:M{M_END}")
dv_status.add(f"I2:I{M_END}"); dv_goals.add(f"E2:F{M_END}")
ma.conditional_formatting.add(f"A2:M{M_END}", FormulaRule(formula=['$I2="Live"'], fill=LIVE_FILL))
ma.conditional_formatting.add(f"K2:K{M_END}", FormulaRule(formula=['AND($K2<>"",$K2<>"Draw")'], fill=GOLD_FILL))
widths(ma, {"A": 11, "B": 13, "C": 16, "D": 16, "E": 8, "F": 8, "G": 16, "H": 16, "I": 9, "J": 16, "K": 16, "L": 11, "M": 16, "N": 10})
ma.column_dimensions.group("N", "N", hidden=True)
ma.freeze_panes = "A2"

# ---------------- Cash-out ----------------
co = wb.create_sheet("Cash-out")
co["A1"] = "Player"
for i in range(5):
    co.cell(row=1, column=2 + i, value=f'=Settings!A{6+i}&" ("&TEXT(Settings!B{6+i},"$0.00")&")"')
co["G1"], co["H1"], co["I1"] = "Cash out", "Paid in", "Net"
style_header(co, 1, 9)
fill_range(co, f"B2:F{P_END}", IN_FILL)
co["B1"].comment = Comment("Count the chips of each colour each player has at the end.", "Game Night")
for r in range(2, P_END + 1):
    co.cell(row=r, column=1, value=f'=Stats!A{r}')
    co.cell(row=r, column=7, value=f'=IF(COUNT(B{r}:F{r})=0,"",B{r}*Settings!$B$6+C{r}*Settings!$B$7+D{r}*Settings!$B$8+E{r}*Settings!$B$9+F{r}*Settings!$B$10)').number_format = MONEY0
    co.cell(row=r, column=8, value=f'=IF(A{r}="","",Stats!X{r})').number_format = MONEY0
    co.cell(row=r, column=9, value=f'=IF(G{r}="","",G{r}-H{r})').number_format = '+$#,##0.00;-$#,##0.00;$0.00'
co.conditional_formatting.add(f"I2:I{P_END}", FormulaRule(formula=['AND(ISNUMBER($I2),$I2>0)'], font=Font(color="1D7A4C", bold=True)))
co.conditional_formatting.add(f"I2:I{P_END}", FormulaRule(formula=['AND(ISNUMBER($I2),$I2<0)'], font=Font(color="BF3527", bold=True)))
co["K1"] = "Check"; style_header(co, 1, 2, start=11)
co["K2"], co["L2"] = "Total paid in", f"=SUM(Stats!X2:X{P_END})"
co["K3"], co["L3"] = "Total cashed out", f"=SUM(G2:G{P_END})"
co["K4"], co["L4"] = "Difference", "=L3-L2"
co["K5"] = f'=IF(COUNT(G2:G{P_END})=0,"Enter chip counts to check.",IF(ABS(L4)<0.005,"Balanced.",IF(L4>0,"Too much cashed out. Recount.","Money is missing. Recount.")))'
for c in ("L2", "L3", "L4"):
    co[c].number_format = MONEY0
widths(co, {"A": 20, "B": 13, "C": 13, "D": 13, "E": 14, "F": 13, "G": 12, "H": 12, "I": 12, "J": 3, "K": 18, "L": 14})
co.freeze_panes = "B2"

# ---------------- Stats (helper calculations) ----------------
sx = wb.create_sheet("Stats")
sh = ["Name", "Group", "FIFA", "P", "W", "D", "L", "GF", "GA", "GD", "Pts", "Overall sort key",
      "Group P", "Group W", "Group D", "Group L", "Group GF", "Group GA", "Group GD", "Group Pts", "Group no.", "Group sort key",
      "Buy-ins", "Paid in", "Cash out", "Net", "Poker sort key"]
for i, h in enumerate(sh, start=1):
    sx.cell(row=1, column=i, value=h)
style_header(sx, 1, len(sh))
m = lambda col: f"Matches!${col}:${col}"
mA, mC, mE, mF, mH, mI, mK, mL = (m(c) for c in "ACEFHIKL")
bB, bC = "'Buy-ins'!$B:$B", "'Buy-ins'!$C:$C"
FIN = f'{mI},"Final",{mA},"<>Friendly"'
for r in range(2, P_END + 1):
    A = f"$A{r}"
    f = {
        "A": f'=IF(TRIM(Players!A{r})="","",TRIM(Players!A{r}))',
        "B": f'=IF($A{r}="","",Players!B{r}&"")',
        "C": f'=Players!E{r}&""',
        "D": f'=IF({A}="",0,COUNTIFS({mC},{A},{FIN})+COUNTIFS({mH},{A},{FIN}))',
        "E": f'=IF({A}="",0,COUNTIFS({mK},{A},{mA},"<>Friendly"))',
        "F": f'=IF({A}="",0,COUNTIFS({mC},{A},{mK},"Draw",{mA},"<>Friendly")+COUNTIFS({mH},{A},{mK},"Draw",{mA},"<>Friendly"))',
        "G": f'=D{r}-E{r}-F{r}',
        "H": f'=IF({A}="",0,SUMIFS({mE},{mC},{A},{FIN})+SUMIFS({mF},{mH},{A},{FIN}))',
        "I": f'=IF({A}="",0,SUMIFS({mF},{mC},{A},{FIN})+SUMIFS({mE},{mH},{A},{FIN}))',
        "J": f'=H{r}-I{r}',
        "K": f'=3*E{r}+F{r}',
        "L": f'=IF(AND({A}<>"",OR($C{r}="Yes",D{r}>0)),K{r}*1000000+(MAX(-499,MIN(499,J{r}))+500)*1000+MIN(H{r},99)*10+(50-ROW())/10,-1)',
        "M": f'=IF($B{r}="",0,COUNTIFS({mC},{A},{mI},"Final",{mA},"Group",{mL},$B{r})+COUNTIFS({mH},{A},{mI},"Final",{mA},"Group",{mL},$B{r}))',
        "N": f'=IF($B{r}="",0,COUNTIFS({mK},{A},{mA},"Group",{mL},$B{r}))',
        "O": f'=IF($B{r}="",0,COUNTIFS({mC},{A},{mK},"Draw",{mA},"Group",{mL},$B{r})+COUNTIFS({mH},{A},{mK},"Draw",{mA},"Group",{mL},$B{r}))',
        "P": f'=M{r}-N{r}-O{r}',
        "Q": f'=IF($B{r}="",0,SUMIFS({mE},{mC},{A},{mI},"Final",{mA},"Group",{mL},$B{r})+SUMIFS({mF},{mH},{A},{mI},"Final",{mA},"Group",{mL},$B{r}))',
        "R": f'=IF($B{r}="",0,SUMIFS({mF},{mC},{A},{mI},"Final",{mA},"Group",{mL},$B{r})+SUMIFS({mE},{mH},{A},{mI},"Final",{mA},"Group",{mL},$B{r}))',
        "S": f'=Q{r}-R{r}',
        "T": f'=3*N{r}+O{r}',
        "U": f'=IF($B{r}="",0,IFERROR(MATCH($B{r},{GROUPS},0),0))',
        "V": f'=IF(OR({A}="",U{r}=0),-1,U{r}*1000000000+T{r}*1000000+(MAX(-499,MIN(499,S{r}))+500)*1000+MIN(Q{r},99)*10+(50-ROW())/10)',
        "W": f'=IF({A}="",0,COUNTIF({bB},{A}))',
        "X": f'=IF({A}="",0,SUMIF({bB},{A},{bC})+Settings!$B$3*COUNTIFS({bB},{A},{bC},""))',
        "Y": f"=IF({A}=\"\",\"\",'Cash-out'!G{r})",
        "Z": f'=IF(OR({A}="",Y{r}=""),"",Y{r}-X{r})',
        "AA": f'=IF(X{r}=0,-1,(IF(Z{r}="",0,Z{r})+100000)*1000000+ROUND(X{r}*10,0)*10+(50-ROW())/10)',
    }
    for col, formula in f.items():
        sx[f"{col}{r}"] = formula
for col in ("X", "Y", "Z"):
    for r in range(2, P_END + 1):
        sx[f"{col}{r}"].number_format = MONEY0
sx["AC1"] = "Helper calculations for the Leaderboard. Nothing to fill in here."
sx["AC1"].font = Font(name=FONT, italic=True, color="5B6662")
for col in "ABC":
    sx.column_dimensions[col].width = 16
sx.freeze_panes = "B2"

# ---------------- Leaderboard ----------------
lb = wb.create_sheet("Leaderboard", 1)
lb["A1"] = "Leaderboard"; lb["A1"].font = Font(name=FONT, size=18, bold=True, color="1F4E3D")
row = 3

def section(title):
    global row
    lb.cell(row=row, column=1, value=title)
    for c in range(1, 12):
        lb.cell(row=row, column=c).fill = SEC_FILL
    lb.cell(row=row, column=1).font = Font(name=FONT, size=13, bold=True)
    row += 1

# Live now
section("Live now")
for k in range(1, 5):
    r = row
    lb[f"M{r}"] = f'=IFERROR(SMALL(Matches!$N:$N,{k}),"")'
    I = lambda col: f"INDEX(Matches!${col}:${col},$M{r})"
    t1 = f'IF({I("D")}="",""," ("&{I("D")}&")")'
    t2 = f'IF({I("G")}="",""," ("&{I("G")}&")")'
    lb[f"A{r}"] = (f'=IF($M{r}="",' + ('"No matches live right now."' if k == 1 else '""') +
                   f',{I("C")}&{t1}&"   "&N({I("E")})&" - "&N({I("F")})&"   "&{I("H")}&{t2}&"   ["&{I("A")}&"]")')
    lb[f"A{r}"].font = Font(name=FONT, size=12, bold=True)
    row += 1
row += 1

STAT_HDR = ["Rank", "Player", "Group", "P", "W", "D", "L", "GF", "GA", "GD", "Pts"]

def table(title_formula, key_formula, valid_formula, cols, nrows, hdr=STAT_HDR):
    """cols: list of Stats column letters for columns C.. (after Rank, Player)."""
    global row
    lb.cell(row=row, column=1, value=title_formula)
    for c in range(1, 12):
        lb.cell(row=row, column=c).fill = SEC_FILL
    lb.cell(row=row, column=1).font = Font(name=FONT, size=13, bold=True)
    row += 1
    for i, h in enumerate(hdr, start=1):
        lb.cell(row=row, column=i, value=h)
    style_header(lb, row, len(hdr))
    row += 1
    first = row
    for k in range(1, nrows + 1):
        r = row
        lb[f"M{r}"] = key_formula(k)
        lb[f"N{r}"] = f'=IF({valid_formula(r)},MATCH($M{r},{KEYRANGE[0]},0),"")'
        lb[f"A{r}"] = f'=IF($N{r}="","",{k})'
        lb[f"B{r}"] = f'=IF($N{r}="","",INDEX(Stats!$A$2:$A${P_END},$N{r}))'
        for j, sc in enumerate(cols):
            lb.cell(row=r, column=3 + j, value=f'=IF($N{r}="","",INDEX(Stats!${sc}$2:${sc}${P_END},$N{r}))')
        lb[f"B{r}"].font = Font(name=FONT, bold=True)
        for c in range(1, len(hdr) + 1):
            lb.cell(row=r, column=c).border = BORDER
        row += 1
    lb.conditional_formatting.add(f"A{first}:{chr(64+len(hdr))}{first}", FormulaRule(formula=[f'$N{first}<>""'], fill=GOLD_FILL))
    return first

# Group tables
KEYRANGE = [f"Stats!$V$2:$V${P_END}"]
for g in range(1, 5):
    table(f'=IF(Settings!$A${14+g}="","Group {g} (add the name on the Settings tab)",Settings!$A${14+g})',
          lambda k, g=g: f'=IFERROR(LARGE({KEYRANGE[0]},COUNTIF({KEYRANGE[0]},">="&{g+1}*1000000000)+{k}),-1)',
          lambda r, g=g: f'$M{r}>={g}*1000000000',
          ["B", "M", "N", "O", "P", "Q", "R", "S", "T"], 6)
    row += 1

# Overall FIFA
KEYRANGE[0] = f"Stats!$L$2:$L${P_END}"
table("FIFA overall (group and knockout matches)",
      lambda k: f'=IFERROR(LARGE({KEYRANGE[0]},{k}),-1)',
      lambda r: f'$M{r}>=0',
      ["B", "D", "E", "F", "G", "H", "I", "J", "K"], 14)
row += 1

# Poker
KEYRANGE[0] = f"Stats!$AA$2:$AA${P_END}"
lb.cell(row=row, column=1, value="Pot")
lb.cell(row=row, column=2, value=f"=SUM(Stats!X2:X{P_END})").number_format = MONEY0
lb.cell(row=row, column=3, value=f'=SUM(Stats!W2:W{P_END})&" buy-ins"')
lb.cell(row=row, column=1).font = lb.cell(row=row, column=2).font = Font(name=FONT, bold=True, size=12)
pk_title_row = row + 1
row += 1
first = table("Poker (ranked by net once cash-outs are in)",
              lambda k: f'=IFERROR(LARGE({KEYRANGE[0]},{k}),-1)',
              lambda r: f'$M{r}>=0',
              ["W", "X", "Y", "Z"], 13, hdr=["Rank", "Player", "Buy-ins", "Paid in", "Cash out", "Net"])
for r in range(first, first + 13):
    for c in ("D", "E"):
        lb[f"{c}{r}"].number_format = MONEY0
    lb[f"F{r}"].number_format = '+$#,##0.00;-$#,##0.00;$0.00'
lb.conditional_formatting.add(f"F{first}:F{first+12}", FormulaRule(formula=[f'AND(ISNUMBER($F{first}),$F{first}>0)'], font=Font(color="1D7A4C", bold=True)))
lb.conditional_formatting.add(f"F{first}:F{first+12}", FormulaRule(formula=[f'AND(ISNUMBER($F{first}),$F{first}<0)'], font=Font(color="BF3527", bold=True)))
lb["M2"], lb["N2"] = "sort key", "row"
widths(lb, {"A": 7, "B": 20, "C": 12, "D": 9, "E": 9, "F": 9, "G": 6, "H": 6, "I": 6, "J": 6, "K": 6, "L": 3, "M": 12, "N": 6})
lb.column_dimensions.group("M", "N", hidden=True)
# the Live now text spans across: keep column A wide enough via no wrap
for r in range(4, 8):
    lb[f"A{r}"].alignment = Alignment(wrap_text=False)

# Fonts everywhere
for ws in wb.worksheets:
    for rr in ws.iter_rows():
        for cell in rr:
            f0 = cell.font
            if f0.name != FONT:
                cell.font = Font(name=FONT, size=f0.size, bold=f0.bold, italic=f0.italic, color=f0.color)

# Order tabs
order = ["Start here", "Leaderboard", "Players", "Buy-ins", "Matches", "Cash-out", "Settings", "Stats"]
wb._sheets = [wb[n] for n in order]

if SAMPLE:
    se["A15"], se["A16"] = "Group A", "Group B"
    names = [("Ana", "Group A", "Real Madrid", "Yes", "Yes"), ("Ben", "Group A", "Arsenal", "Yes", "Yes"),
             ("Cy", "Group A", "", "No", "Yes"), ("Dee", "Group B", "PSG", "Yes", "Yes"), ("Eli", "Group B", "", "No", "Yes"),
             ("Fay", "Group B", "", "Yes", "No")]
    for i, rowv in enumerate(names, start=2):
        for j, v in enumerate(rowv, start=1):
            pl.cell(row=i, column=j, value=v)
    for i, (p, a) in enumerate([("Ana", None), ("Ana", None), ("Ben", None), ("Dee", 10), ("Fay", None)], start=2):
        bi.cell(row=i, column=2, value=p); bi.cell(row=i, column=3, value=a)
    ms = [("Group", None, "Ana", "Real Madrid", 2, 0, "Arsenal", "Ben", "Final", None),
          ("Group", None, "Ben", "", 1, 1, "", "Cy", "Final", None),
          ("Group", None, "Dee", "PSG", 3, 1, "", "Eli", "Final", None),
          ("Knockout", "Final", "Ana", "", 1, 1, "", "Dee", "Final", "Dee"),
          ("Friendly", None, "Cy", "", 5, 0, "", "Eli", "Final", None),
          ("Group", None, "Ana", "", 0, 1, "", "Cy", "Live", None)]
    for i, rowv in enumerate(ms, start=2):
        for j, v in enumerate(rowv, start=1):
            if v is not None:
                ma.cell(row=i, column=j, value=v)
    # cash-out rows align with players: Ana, Ben, Cy, Dee, Eli, Fay
    for r, counts in [(2, [10, 10, 4, 10, 1]), (3, [0, 0, 0, 0, 0]), (5, [20, 20, 8, 20, 2]), (7, [10, 10, 4, 10, 0])]:
        for j, v in enumerate(counts):
            co.cell(row=r, column=2 + j, value=v)

wb.save(out)
print("saved", out)
